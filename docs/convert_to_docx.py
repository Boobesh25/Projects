"""Convert the Markdown study guide to a properly formatted DOCX."""
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

# Set default font
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

with open("docs/AI_ML_RAG_Study_Guide.md", "r", encoding="utf-8") as f:
    content = f.read()

lines = content.split("\n")
i = 0
in_code_block = False
code_lines = []

while i < len(lines):
    line = lines[i]

    # Code blocks
    if line.strip().startswith("```"):
        if in_code_block:
            # End code block — add as formatted paragraph
            code_text = "\n".join(code_lines)
            p = doc.add_paragraph()
            run = p.add_run(code_text)
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            p.paragraph_format.left_indent = Cm(1)
            code_lines = []
            in_code_block = False
        else:
            in_code_block = True
        i += 1
        continue

    if in_code_block:
        code_lines.append(line)
        i += 1
        continue

    # Tables (collect all | rows)
    if line.strip().startswith("|") and "|" in line[1:]:
        table_rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            row_line = lines[i].strip()
            # Skip separator rows
            if re.match(r"^\|[\s\-:|]+\|$", row_line):
                i += 1
                continue
            cells = [c.strip() for c in row_line.split("|")[1:-1]]
            table_rows.append(cells)
            i += 1

        if table_rows:
            num_cols = max(len(row) for row in table_rows)
            table = doc.add_table(rows=len(table_rows), cols=num_cols)
            table.style = 'Table Grid'

            for row_idx, row_data in enumerate(table_rows):
                for col_idx, cell_text in enumerate(row_data):
                    if col_idx < num_cols:
                        cell = table.cell(row_idx, col_idx)
                        # Clean markdown bold
                        clean_text = cell_text.replace("**", "")
                        cell.text = clean_text
                        # Bold header row
                        if row_idx == 0:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.bold = True

            doc.add_paragraph()  # Space after table
        continue

    # Headings
    if line.startswith("# "):
        doc.add_heading(line[2:].strip(), level=0)
        i += 1
        continue
    elif line.startswith("## "):
        doc.add_heading(line[3:].strip(), level=1)
        i += 1
        continue
    elif line.startswith("### "):
        doc.add_heading(line[4:].strip(), level=2)
        i += 1
        continue
    elif line.startswith("#### "):
        doc.add_heading(line[5:].strip(), level=3)
        i += 1
        continue

    # Horizontal rule
    if line.strip() == "---":
        p = doc.add_paragraph()
        p.add_run("─" * 60).font.color.rgb = RGBColor(180, 180, 180)
        i += 1
        continue

    # Blockquotes (interview answers)
    if line.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.5)
        run = p.add_run(line[2:])
        run.italic = True
        run.font.color.rgb = RGBColor(60, 60, 120)
        i += 1
        continue

    # Bullet points
    if line.startswith("- ") or line.startswith("* "):
        text = line[2:].replace("**", "")
        doc.add_paragraph(text, style='List Bullet')
        i += 1
        continue

    # Numbered items
    if re.match(r"^\d+\.\s", line):
        text = re.sub(r"^\d+\.\s", "", line).replace("**", "")
        doc.add_paragraph(text, style='List Number')
        i += 1
        continue

    # Empty lines
    if not line.strip():
        i += 1
        continue

    # Bold text handling
    if "**" in line:
        p = doc.add_paragraph()
        parts = re.split(r"(\*\*.*?\*\*)", line)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                p.add_run(part)
        i += 1
        continue

    # Regular text
    doc.add_paragraph(line)
    i += 1

doc.save("docs/AI_ML_RAG_Study_Guide.docx")
print("✅ Saved: docs/AI_ML_RAG_Study_Guide.docx")
